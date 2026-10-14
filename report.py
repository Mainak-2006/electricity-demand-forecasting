"""BijliPredict AI - assemble an Excel report from pipeline outputs.

Builds a single multi-sheet workbook from the files produced by
`predict.py` and `evaluate.py`:

    Summary             bill + projection headline numbers
    Forecast            the 7-day forecast (kWh + cost)
    Projection          the 90-day rolling projection
    Evaluation          MAE / RMSE / WAPE vs baselines
    Tariff              the tariff used, per-day costs

Usage:
    python report.py                          # from default outputs/
    python report.py --outputs-dir outputs --out outputs/bijli_report.xlsx
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_OUTPUTS = "outputs"
DEFAULT_XLSX = "outputs/bijli_report.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _read_json(path: Path) -> dict:
    if not path.exists():
        return {}
    return json.loads(path.read_text())


def _read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path)


def _summary_frame(cost: dict, evaluation: dict) -> pd.DataFrame:
    proj = cost.get("projection", {})
    monthly = proj.get("monthly_30d", {})
    quarterly = proj.get("three_month_90d", {})
    model_m = evaluation.get("model_metrics", {})
    base_m = evaluation.get("baseline_metrics", {}).get("seasonal_naive_p7", {})

    rows = [
        ("Generated", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Model", evaluation.get("model", "gru").upper()),
        ("Tariff", cost.get("tariff", "")),
        ("", ""),
        ("7-day forecast (kWh)", round(cost.get("total_kwh", 0.0), 2)),
        ("7-day estimated bill (INR)", cost.get("estimated_bill_inr", 0.0)),
        ("Monthly 30d (kWh)", round(monthly.get("total_kwh", 0.0), 2)),
        ("Monthly 30d bill (INR)", monthly.get("estimated_bill_inr", 0.0)),
        ("3-Month 90d (kWh)", round(quarterly.get("total_kwh", 0.0), 2)),
        ("3-Month 90d bill (INR)", quarterly.get("estimated_bill_inr", 0.0)),
        ("", ""),
        ("Test MAE (model)", round(model_m.get("mae", 0.0), 3)),
        ("Test RMSE (model)", round(model_m.get("rmse", 0.0), 3)),
        ("Test WAPE % (model)", round(model_m.get("wape", 0.0), 2)),
        ("Test WAPE % (seasonal-naive)", round(base_m.get("wape", 0.0), 2)),
    ]
    return pd.DataFrame(rows, columns=["Item", "Value"])


def _evaluation_frame(evaluation: dict) -> pd.DataFrame:
    model_m = evaluation.get("model_metrics", {})
    baselines = evaluation.get("baseline_metrics", {})
    metrics = ["mae", "rmse", "wape", "mape"]
    data = {
        "metric": [m.upper() for m in metrics],
        evaluation.get("model", "model").upper(): [model_m.get(m) for m in metrics],
        "seasonal_naive_p7": [baselines.get("seasonal_naive_p7", {}).get(m) for m in metrics],
        "persistence": [baselines.get("persistence", {}).get(m) for m in metrics],
    }
    return pd.DataFrame(data)


def _tariff_frame(cost: dict) -> pd.DataFrame:
    rows = [("Tariff", cost.get("tariff", ""))]
    daily = cost.get("daily", [])
    for d in daily:
        rows.append((d.get("date"), d.get("cost_inr")))
    return pd.DataFrame(rows, columns=["Date / Info", "Daily cost (INR)"])


def _autofit(writer: pd.ExcelWriter) -> None:
    for ws in writer.book.worksheets:
        for col_cells in ws.columns:
            length = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(
                40, max(12, length + 2)
            )
        for cell in ws[1]:
            if cell.value is not None:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"


def build_excel(
    outputs_dir: str | Path = DEFAULT_OUTPUTS,
    out_path: str | Path = DEFAULT_XLSX,
) -> Path:
    """Read pipeline outputs and write a multi-sheet Excel workbook."""
    outputs_dir = Path(outputs_dir)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    cost = _read_json(outputs_dir / "cost_summary.json")
    evaluation = _read_json(outputs_dir / "evaluation.json")
    forecast = _read_csv(outputs_dir / "predictions.csv")
    projection = _read_csv(outputs_dir / "projection_90d.csv")

    sheets: dict[str, pd.DataFrame] = {
        "Summary": _summary_frame(cost, evaluation),
        "Forecast": forecast,
        "Projection": projection,
        "Evaluation": _evaluation_frame(evaluation),
        "Tariff": _tariff_frame(cost),
    }

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            (frame if not frame.empty else pd.DataFrame({"info": ["no data"]})).to_excel(
                writer, sheet_name=name, index=False
            )
        _autofit(writer)
    return out_path


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Build the Excel report from outputs.")
    p.add_argument("--outputs-dir", default=DEFAULT_OUTPUTS)
    p.add_argument("--out", default=DEFAULT_XLSX)
    return p


def main(args: argparse.Namespace | None = None) -> None:
    if args is None:
        args = build_parser().parse_args()
    path = build_excel(args.outputs_dir, args.out)
    print(f"Excel report saved: {path}")


if __name__ == "__main__":
    main()
